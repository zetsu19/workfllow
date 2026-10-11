import os, tempfile
from datetime import date
from receipt_parser import parse_items, find_date
from excel import append_items

TXT = """Огноо 2026-10-06
Хуванцар таваг 10 7,500 75,000
Хар таваг зуйван
10 x 4,800 48,000
Шар аяга тагтай 5 29,900 149,500
Өндөг хутгагч 6,000
Mixed seafood 36,900
Нийт дүн 315,400
НӨАТ 10%
"""
def test_all():
    it = parse_items(TXT)
    got = [(i.name, i.qty, i.price, i.total) for i in it]
    assert got == [("Хуванцар таваг", 10, 7500, 75000), ("Хар таваг зуйван", 10, 4800, 48000),
                   ("Шар аяга тагтай", 5, 29900, 149500), ("Өндөг хутгагч", 1, 6000, 6000),
                   ("Mixed seafood", 1, 36900, 36900)], got
    assert find_date(TXT) == date(2026, 10, 6)
    p = os.path.join(tempfile.mkdtemp(), "a.xlsx")
    append_items(p, date(2026, 10, 6), it); append_items(p, date(2026, 10, 7), it[:1])
    from openpyxl import load_workbook
    ws = load_workbook(p).active
    assert ws.max_row == 7 and ws["F2"].value == "=D2*E2"
if __name__ == "__main__":
    test_all(); print("ok")
