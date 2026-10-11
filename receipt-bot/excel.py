import os
from datetime import date
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

HEADERS = ["Огноо", "нэр", "нэмжих нэгж", "тоо", "нэгжийн үнэ", "Нийт дүн", "Дэлгүүр", "Падааны нийт"]
WIDTHS = (13, 38, 14, 8, 14, 14, 24, 16)
HEAD_FILL = PatternFill("solid", fgColor="B4A7D6")
NEW_FILL = PatternFill("solid", fgColor="B7DEE8")  # шинэ (сүүлд нэмсэн) мөрүүд
NCOL = len(HEADERS)


def _style_header(ws):
    for c, w in zip("ABCDEFGH", WIDTHS):
        ws.column_dimensions[c].width = w
    for i, h in enumerate(HEADERS, 1):
        cell = ws.cell(1, i, h)
        cell.fill, cell.font = HEAD_FILL, Font(bold=True)


def _new():
    wb = Workbook()
    ws = wb.active
    ws.title = "Худалдан авалт"
    _style_header(ws)
    return wb


def append_items(path, d: date, items, store=None):
    """Мөрүүдийг нэмнэ. Нийт дүнг (тоо×үнэ) утгаар бичнэ (томьёо биш: preview/Protected View-д хоосон харагддаг).
    Падааны нийт дүн: тухайн падааны эхний мөрийн хамгийн ард баганад. -> эхний мөрийн дугаар"""
    wb = load_workbook(path) if os.path.exists(path) else _new()
    ws = wb.active
    if ws.cell(1, 7).value is None:  # хуучин 6 баганатай файл
        _style_header(ws)
    for t in list(ws.tables):
        del ws.tables[t]
    for row in ws.iter_rows(min_row=2):  # өмнөх шинэ тэмдэглэгээг арилгана
        for c in row:
            c.fill = PatternFill(fill_type=None)
    first = ws.max_row + 1
    for it in items:
        r = ws.max_row + 1
        ws.append([d, it.name, None, it.qty, it.price, it.total, store])
        ws.cell(r, 1).number_format = "d/m/yyyy"
        for col in (5, 6):
            ws.cell(r, col).number_format = "#,##0"
        ws.cell(r, 6).font = Font(bold=True)
        for col in range(1, NCOL + 1):
            ws.cell(r, col).fill = NEW_FILL
    _batch_total(ws, first, len(items))
    ws.auto_filter.ref = f"A1:H{ws.max_row}"
    ws.freeze_panes = "A2"
    wb.save(path)
    return first


def _batch_total(ws, start, count):
    total = sum(ws.cell(r, 6).value or 0 for r in range(start, start + count))
    c = ws.cell(start, 8, round(total, 2))
    c.number_format, c.font = "#,##0", Font(bold=True)
    c.fill = NEW_FILL


def update_row(path, row, name=None, qty=None, price=None, batch=None):
    """-> хуучин нэр (эсвэл мөр олдохгүй бол None). Нийт дүн болон падааны нийтийг дахин бодно.
    batch=(эхний мөр, тоо) өгвөл падааны нийтийг шинэчилнэ."""
    wb = load_workbook(path)
    ws = wb.active
    if row < 2 or row > ws.max_row:
        return None
    old = ws.cell(row, 2).value
    if name is not None:
        ws.cell(row, 2).value = name
    if qty is not None:
        ws.cell(row, 4).value = qty
    if price is not None:
        ws.cell(row, 5).value = price
    ws.cell(row, 6).value = round((ws.cell(row, 4).value or 0) * (ws.cell(row, 5).value or 0), 2)
    if batch:
        _batch_total(ws, *batch)
    wb.save(path)
    return old
