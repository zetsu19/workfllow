"""AI давхаргыг жинхэнэ API дуудахгүйгээр (mock сервер) шалгана."""
import json, os, tempfile, threading
from http.server import BaseHTTPRequestHandler, HTTPServer
import anthropic
import ai_vision
from excel import append_items, update_row
from memory import Memory
from receipt_parser import Item

CALLS = []


def reply(model):
    items = [{"name": "Зөгийн бал", "qty": 1, "unit_price": 45890, "total": 45890},
             {"name": "Үсний будаг", "qty": 2, "unit_price": 31990, "total": 63980}]
    gt = 109870 if model.startswith("claude-sonnet") else 999999  # haiku: таараагүй -> sonnet руу шилжинэ
    return {"id": "msg_1", "type": "message", "role": "assistant", "model": model, "stop_reason": "end_turn",
            "stop_sequence": None, "usage": {"input_tokens": 10, "output_tokens": 10},
            "content": [{"type": "text", "text": json.dumps({"store": "ХМ Наадам центр", "date": "2026-10-11", "grand_total": gt, "items": items})}]}


class H(BaseHTTPRequestHandler):
    def do_POST(self):
        body = json.loads(self.rfile.read(int(self.headers["content-length"])))
        CALLS.append(body)
        out = json.dumps(reply(body["model"])).encode()
        self.send_response(200); self.send_header("content-type", "application/json")
        self.send_header("content-length", str(len(out))); self.end_headers(); self.wfile.write(out)
    def log_message(self, *a): pass


def test():
    srv = HTTPServer(("127.0.0.1", 0), H); threading.Thread(target=srv.serve_forever, daemon=True).start()
    client = anthropic.Anthropic(api_key="x", base_url=f"http://127.0.0.1:{srv.server_port}")
    import cv2, numpy as np
    img = cv2.imencode(".png", np.full((300, 200, 3), 255, np.uint8))[1].tobytes()
    res = ai_vision.read_with_ai(img, ["Зөгийн бал"], client=client)
    items, gt, matched, when, model, store = res
    assert store == "ХМ Наадам центр"
    assert [c["model"] for c in CALLS] == ["claude-haiku-5-5", "claude-sonnet-5-5"], CALLS
    assert matched and model == "claude-sonnet-5-5" and len(items) == 2 and str(when) == "2026-10-11"
    b = CALLS[0]
    assert b["messages"][0]["content"][0]["type"] == "image" and b["output_config"]["effort"] == "low"
    assert "Зөгийн бал" in b["messages"][0]["content"][1]["text"] and b["output_config"]["format"]["type"] == "json_schema"
    # санах ой + засвар
    d = tempfile.mkdtemp(); m = Memory(os.path.join(d, "m.json"))
    p = os.path.join(d, "a.xlsx"); first = append_items(p, when, [Item("Valhalla Hrb", 3, 66000, 198000), Item("X", 1.07, 14780, 15815)], "ХМ")
    old = update_row(p, first, "Valhalla Herb liqueur 35%", 2, 66000, (first, 2)); m.learn(old, "Valhalla Herb liqueur 35%"); m.save()
    it = [Item("valhalla  hrb", 1, 1, 1)]; Memory(os.path.join(d, "m.json")).apply(it)
    assert it[0].name == "Valhalla Herb liqueur 35%"
    from openpyxl import load_workbook
    ws = load_workbook(p).active
    assert [c.value for c in ws[1]][-2:] == ["Дэлгүүр", "Падааны нийт"]
    assert ws["F2"].value == 132000 and ws["G2"].value == "ХМ" and ws["H2"].value == 132000 + 15815 and ws["H3"].value is None, [c.value for c in ws[2]]
    print("ok")

if __name__ == "__main__":
    test()
